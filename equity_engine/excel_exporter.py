import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from .database import get_connection
from .calculator import run_vm_analysis

def generate_analysis_excel(symbol, output_path=None, db_path=None):
    # Ensure analysis is run and up-to-date
    analysis = run_vm_analysis(symbol, db_path=db_path)
    
    conn = get_connection(db_path) if db_path else get_connection()
    comp = conn.execute("SELECT * FROM companies WHERE symbol=?", (symbol,)).fetchone()
    df_ann = pd.read_sql("SELECT * FROM annual_financials WHERE symbol=? ORDER BY fiscal_year ASC", conn, params=(symbol,))
    df_q = pd.read_sql("SELECT * FROM quarterly_financials WHERE symbol=? ORDER BY period ASC", conn, params=(symbol,))
    mkt = conn.execute("SELECT * FROM daily_market_quotes WHERE symbol=? ORDER BY date DESC LIMIT 1", (symbol,)).fetchone()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VM Analysis"

    # Styling definitions
    font_title = Font(name="Calibri", size=11, bold=True)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_italic = Font(name="Calibri", size=9, italic=True)

    fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # soft blue
    fill_score = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # soft green
    fill_subhead = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # light gray

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Row 1
    ws.cell(1, 1, "Shri Sai").font = font_title

    # Row 3: Summary Banner
    ws.cell(3, 1, comp['name'].upper()).font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    
    ws.cell(3, 6, "Quality").font = font_bold
    ws.cell(3, 7, analysis['quality_score']).font = font_bold
    ws.cell(3, 8, analysis['max_quality']).font = font_bold
    ws.cell(3, 9, round(analysis['quality_score'] / analysis['max_quality'], 4)).font = font_bold

    ws.cell(3, 10, "Valuation").font = font_bold
    ws.cell(3, 11, analysis['valuation_score']).font = font_bold
    ws.cell(3, 12, analysis['max_valuation']).font = font_bold
    ws.cell(3, 13, round(analysis['valuation_score'] / analysis['max_valuation'], 4)).font = font_bold
    ws.cell(3, 14, "Combined").font = font_bold

    # Row 5: Table Header
    headers = [
        (1, "No."), (2, "Parameter"), (6, "Criteria"),
        (10, "Actual Score"), (11, "Max Score"), (12, "Comments")
    ]
    for col_idx, h_text in headers:
        cell = ws.cell(5, col_idx, h_text)
        cell.font = font_bold
        cell.fill = fill_header

    # Populate 34 Parameters
    row_cursor = 6
    param_row_map = {}

    for p_num, p in analysis['parameters'].items():
        param_row_map[p_num] = row_cursor
        ws.cell(row_cursor, 1, p_num).font = font_bold
        ws.cell(row_cursor, 2, p['name']).font = font_bold
        ws.cell(row_cursor, 6, p['criteria']).font = font_regular
        
        score_cell = ws.cell(row_cursor, 10, p['score'])
        score_cell.font = font_bold
        score_cell.fill = fill_score
        
        ws.cell(row_cursor, 11, p['max_score']).font = font_regular

        row_cursor += 1

    # Add Data Tables underneath
    row_cursor += 2
    ws.cell(row_cursor, 1, "HISTORICAL FINANCIAL STATEMENTS (RAW DATA)").font = Font(name="Calibri", size=12, bold=True, color="1F497D")
    row_cursor += 1

    # Helper function to write a multi-year table
    year_cols = list(df_ann['period'])
    def write_table(title, rows_data):
        nonlocal row_cursor
        ws.cell(row_cursor, 2, title).font = font_bold
        ws.cell(row_cursor, 2).fill = fill_subhead
        
        # Header dates
        for idx, col_name in enumerate(year_cols):
            c = ws.cell(row_cursor, 5 + idx, col_name)
            c.font = font_bold
            c.fill = fill_subhead
            c.alignment = Alignment(horizontal="right")
        row_cursor += 1

        for label, series in rows_data:
            ws.cell(row_cursor, 2, label).font = font_regular
            for idx, val in enumerate(series):
                c = ws.cell(row_cursor, 5 + idx, round(val, 2) if isinstance(val, (int, float)) else val)
                c.font = font_regular
                c.alignment = Alignment(horizontal="right")
                c.border = thin_border
            row_cursor += 1
        row_cursor += 1

    # 1. P&L Statement
    write_table("Annual Profit & Loss (Rs. Cr)", [
        ("Sales", df_ann['sales'].tolist()),
        ("Expenses", df_ann['expenses'].tolist()),
        ("Operating Profit", df_ann['operating_profit'].tolist()),
        ("OPM (%)", [f"{x:.1f}%" for x in df_ann['opm_pct']]),
        ("Depreciation", df_ann['depreciation'].tolist()),
        ("Interest", df_ann['interest'].tolist()),
        ("Profit before Tax", df_ann['pbt'].tolist()),
        ("Tax Amount", df_ann['tax_amount'].tolist()),
        ("Net Profit (PAT)", df_ann['pat'].tolist()),
        ("NPM (%)", [f"{(p/s*100):.1f}%" if s>0 else "0%" for p, s in zip(df_ann['pat'], df_ann['sales'])])
    ])

    # 2. Balance Sheet
    write_table("Annual Balance Sheet (Rs. Cr)", [
        ("Equity Share Capital", df_ann['equity_capital'].tolist()),
        ("Reserves", df_ann['reserves'].tolist()),
        ("Net Worth", df_ann['net_worth'].tolist()),
        ("Borrowings (Debt)", df_ann['borrowings'].tolist()),
        ("Fixed Assets (Net Block)", df_ann['fixed_assets'].tolist()),
        ("Capital WIP", df_ann['cwip'].tolist()),
        ("Investments", df_ann['investments'].tolist()),
        ("Other Assets", df_ann['other_assets'].tolist()),
        ("Other Liabilities", df_ann['other_liabilities'].tolist())
    ])

    # 3. Cash Flows
    write_table("Cash Flows & Free Cash Flow (Rs. Cr)", [
        ("Cash from Operations (CFO)", df_ann['cfo'].tolist()),
        ("Capex", df_ann['capex'].tolist()),
        ("Free Cash Flow (FCF)", df_ann['fcf'].tolist()),
        ("Dividend Paid", df_ann['dividend_amount'].tolist())
    ])

    # 4. Valuation Ratios
    row_cursor += 1
    ws.cell(row_cursor, 2, "Current Market Valuation").font = font_bold
    ws.cell(row_cursor, 2).fill = fill_subhead
    row_cursor += 1

    val_items = [
        ("Current Market Price (CMP)", f"Rs. {mkt['cmp']}"),
        ("Market Capitalization", f"Rs. {mkt['market_cap']} Cr"),
        ("Stock P/E Ratio", mkt['pe']),
        ("Price to Book (P/B) Ratio", round(mkt['cmp'] / mkt['book_value'], 2) if mkt['book_value'] else "N/A"),
        ("Dividend Yield", f"{mkt['div_yield']}%"),
        ("ROCE", f"{mkt['roce']}%"),
        ("ROE", f"{mkt['roe']}%"),
        ("10-Year G-Sec Benchmark", f"{mkt['gsec_10y']}%")
    ]
    for lbl, v in val_items:
        ws.cell(row_cursor, 2, lbl).font = font_regular
        c = ws.cell(row_cursor, 5, v)
        c.font = font_bold
        c.alignment = Alignment(horizontal="right")
        row_cursor += 1

    # Auto-fit columns
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 5
    ws.column_dimensions['D'].width = 5
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 36
    ws.column_dimensions['G'].width = 10
    ws.column_dimensions['H'].width = 10
    ws.column_dimensions['I'].width = 10
    ws.column_dimensions['J'].width = 14
    ws.column_dimensions['K'].width = 12
    ws.column_dimensions['L'].width = 15

    for c_idx in range(5, 5 + len(year_cols)):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = 14

    if not output_path:
        output_path = f"{symbol}_Analysis.xlsx"

    wb.save(output_path)
    print(f"Successfully generated analysis workbook at: {output_path}")
    return output_path

if __name__ == "__main__":
    import sys
    s = sys.argv[1] if len(sys.argv) > 1 else "JSWSTEEL"
    out = sys.argv[2] if len(sys.argv) > 2 else f"{s}_analysis_generated.xlsx"
    generate_analysis_excel(s, output_path=out)
