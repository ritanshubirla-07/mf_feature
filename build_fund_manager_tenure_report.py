"""
build_fund_manager_tenure_report.py
===================================
Enriches `equity_fund_manager_history.xlsx` (ACE MF fund-manager stint history) with
tenure and 3-year rolling-window performance attribution, and writes
`equity_fund_manager_history_enriched.xlsx`.

Index funds, ETFs and fund-of-funds are excluded from the output: they track an index
rather than being actively managed, so ranking them against active equity peers says
nothing about a manager. Excluding them changes only which ROWS are listed - peer
groups and quartile cut-offs still come from the analytics engine unchanged, exactly
as index.html computes them.

Added columns
-------------
  Tenure                          stint length as "X years and Y months"
  Tenure (Months)                 same value as an integer, so the column sorts
  Manager - Windows Managed       how many of the 18 half-yearly 3Y windows this manager
                                  held the fund for the ENTIRE window (partial windows are
                                  never counted) AND the fund had a ranked return
  Manager - Times in Top 25%      of those windows, how many the fund finished top-quartile
  Manager - Top 25% Hit Rate      the ratio of the two
  Fund - Windows Ranked           every window the FUND has a 3Y return for, i.e. from its
                                  own start date to the present, whoever was managing
  Fund - Times in Top 25%         of those, how many it finished top-quartile
  Fund - Top 25% Hit Rate         the ratio of the two
  Coverage                        why a row has no window figures, when it has none

Active managers (stint still open at the source file's as-of date) are highlighted.
"""

import math
import re
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, ".")
from fm_core import (build_index, canon, is_data_row, load_engine, load_engine_all,
                     load_stints, match_scheme)

SRC = "equity_fund_manager_history.xlsx"
OUT = "equity_fund_manager_history_enriched.xlsx"

BASE_HEADERS = ["Scheme Name", "Inception Date", "Fund Manager", "From Date", "To Date"]
NEW_HEADERS = ["Tenure", "Tenure (Months)",
               "Manager - Windows Managed", "Manager - Times in Top 25%",
               "Manager - Top 25% Hit Rate",
               "Fund - Windows Ranked", "Fund - Times in Top 25%",
               "Fund - Top 25% Hit Rate", "Coverage"]
HEADERS = BASE_HEADERS + NEW_HEADERS
N_BASE = len(BASE_HEADERS)

# Passive vehicles: by name, and by how the NAV master classifies them.
PASSIVE_NAME = re.compile(r"\bETF\b|\bIndex Fund\b|\bBeES\b|\bIndex Plan\b|"
                          r"\bFund of Fund\b|\bFoF\b", re.I)
PASSIVE_SUBS = {
    "Index Funds", "Indexs", "Equity ETF", "Other ETF", "Other ETFs", "Debt ETF",
    "Gold ETF", "Silver ETF", "Hybrid ETF", "ETFs investing overseas",
    "FoF Domestic", "FoF Overseas", "Fund of Funds Scheme (Domestic)",
    "Fund ofs Scheme (Domestic)", "Fund of Funds investing overseas",
}

NAVY = "1F4E79"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
ACTIVE_FILL = PatternFill("solid", fgColor="D9EAD3")       # green = still managing
GROUP_FILL = PatternFill("solid", fgColor="EDF2F8")
THIN = Side(style="thin", color="BFCDDC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "yyyy-mm-dd"


# ------------------------------------------------------------------ tenure
def tenure_months(frm, to) -> int:
    m = (to.year - frm.year) * 12 + (to.month - frm.month)
    if to.day < frm.day:
        m -= 1
    return max(m, 0)


def tenure_text(months: int) -> str:
    y, m = divmod(months, 12)
    return f"{y} {'year' if y == 1 else 'years'} and {m} {'month' if m == 1 else 'months'}"


# ------------------------------------------------- windows & top-quartile sets
def build_windows(engine, funds):
    import datetime as dt

    starts = [dt.datetime.strptime(s, "%Y-%m-%d") for s in engine["window_starts_iso"]]
    labels = engine["window_dates"]
    windows = [(s, s.replace(year=s.year + 3), f"{labels[i][0]} to {labels[i][1]}")
               for i, s in enumerate(starts)]

    by_sub = defaultdict(list)
    for f in funds.values():
        by_sub[f["sub"]].append(f)

    top25 = []
    for wi in range(len(windows)):
        codes = set()
        for lst in by_sub.values():
            # stable sort on CAGR alone, preserving engine order on ties -- this is what
            # index.html's `.sort((a, b) => b.cagr - a.cagr)` does (JS sort is stable)
            ranked = sorted((f for f in lst if has_cagr(f, wi)),
                            key=lambda f: f["cagrs"][wi], reverse=True)
            if ranked:
                # Math.round semantics (half UP), matching the dashboard's JS
                k = max(1, math.floor(len(ranked) * 0.25 + 0.5))
                codes.update(f["code"] for f in ranked[:k])
        top25.append(codes)
    return windows, top25


def has_cagr(fund, wi) -> bool:
    c = fund.get("cagrs")
    return bool(c) and wi < len(c) and c[wi] is not None


# ------------------------------------------------------- passive / coverage
def engine_subs(scheme, full_index):
    return {h["sub"] for h in full_index.get("".join(sorted(canon(scheme))), [])}


def is_passive(scheme, full_index) -> bool:
    if PASSIVE_NAME.search(scheme):
        return True
    subs = engine_subs(scheme, full_index)
    return bool(subs) and subs <= PASSIVE_SUBS


def coverage_reason(scheme, full_index):
    hits = full_index.get("".join(sorted(canon(scheme))), [])
    if hits:
        plans = {h["plan"] for h in hits}
        opts = {h["option"] for h in hits}
        subs = sorted({h["sub"] for h in hits})
        if "Regular" not in plans:
            return "No Regular plan in NAV master"
        if "Growth" not in opts:
            return "Growth option not labelled in NAV master"
        return f"Not in Equity group (classified: {subs[0]})"
    if re.search(r"long[- ]short|\bSIF\b", scheme, re.I):
        return "Long-short / SIF - no 3Y rolling history"
    return "Not present in analytics engine data"


def main():
    engine, funds = load_engine()
    exact, cands = build_index(funds)
    windows, top25 = build_windows(engine, funds)

    full_index = defaultdict(list)
    for f in load_engine_all():
        full_index["".join(sorted(canon(f["name"], is_engine=True)))].append(f)

    rows = load_stints()
    data_rows = [r for r in rows if is_data_row(r)]
    if not data_rows:
        sys.exit("no data rows parsed from " + SRC)
    asof = max(r["to"] for r in data_rows)

    # ---- decide, once per scheme, whether it stays and what it maps to
    schemes = {r["scheme"].strip() for r in data_rows}
    passive, match_of, reason_of = set(), {}, {}
    for s in schemes:
        if is_passive(s, full_index):
            passive.add(s)
            continue
        code, _score, _m = match_scheme(s, funds, exact, cands)
        match_of[s] = code
        if code is None:
            reason_of[s] = coverage_reason(s, full_index)

    kept = [r for r in data_rows if r["scheme"].strip() not in passive]
    dropped_rows = len(data_rows) - len(kept)
    print(f"source stint rows={len(data_rows)}  as-of={asof:%Y-%m-%d}")
    print(f"passive schemes excluded={len(passive)} ({dropped_rows:,} rows)  "
          f"kept schemes={len(schemes) - len(passive)} ({len(kept):,} rows)")

    # ---- fund-level all-time figures, computed once per engine fund
    fund_alltime = {}
    for code, f in funds.items():
        ranked = [wi for wi in range(len(windows)) if has_cagr(f, wi)]
        fund_alltime[code] = (len(ranked), sum(1 for wi in ranked if code in top25[wi]))

    # ---- write a fresh sheet, preserving the source's grouped layout
    wb = Workbook()
    ws = wb.active
    ws.title = "Fund Manager History"

    t = ws.cell(row=1, column=1, value=">>Fund Manager History  (index funds, ETFs and "
                                      "fund-of-funds excluded)")
    t.font = Font(bold=True, size=12, color=NAVY)
    for j, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER

    # group the kept stints under their scheme, in source order
    order, groups = [], defaultdict(list)
    for r in kept:
        s = r["scheme"].strip()
        if s not in groups:
            order.append(s)
        groups[s].append(r)

    stats = Counter()
    active_rows = 0
    ri = 3
    for scheme in order:
        ri += 1
        g = ws.cell(row=ri, column=1, value=f"Scheme Name: {scheme} ")
        g.font = Font(bold=True, size=10, color=NAVY)
        for j in range(1, len(HEADERS) + 1):
            ws.cell(row=ri, column=j).fill = GROUP_FILL

        code = match_of.get(scheme)
        f_ranked, f_top = fund_alltime.get(code, (None, None)) if code else (None, None)

        for r in groups[scheme]:
            ri += 1
            frm, to = r["from"], r["to"]
            months = tenure_months(frm, to)
            is_active = (to == asof)

            managed = top = None
            if code is not None:
                f = funds[code]
                managed = top = 0
                for wi, (w_start, w_end, _lbl) in enumerate(windows):
                    if frm <= w_start and to >= w_end and has_cagr(f, wi):
                        managed += 1
                        if code in top25[wi]:
                            top += 1
                stats["matched rows"] += 1
                if managed:
                    stats["rows with >=1 full window"] += 1
            else:
                stats["unmatched rows"] += 1

            vals = [
                scheme, r["incep"], r["fm"], frm, to,
                tenure_text(months), months,
                managed, top,
                (round(top / managed, 4) if (code is not None and managed) else None),
                f_ranked, f_top,
                (round(f_top / f_ranked, 4) if (code is not None and f_ranked) else None),
                None if code is not None else reason_of.get(scheme, "Not matched"),
            ]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(row=ri, column=j, value=v)
                c.border = BORDER
                c.font = Font(size=10, bold=is_active)
                if j in (2, 4, 5):
                    c.number_format = DATE_FMT
                    c.alignment = Alignment(horizontal="center")
                elif j in (1, 3, 6):
                    c.alignment = Alignment(horizontal="left")
                elif j in (10, 13):
                    c.number_format = "0.0%"
                    c.alignment = Alignment(horizontal="center")
                elif j == len(HEADERS):
                    c.alignment = Alignment(horizontal="left")
                    c.font = Font(size=9, italic=True, color="7F7F7F")
                else:
                    c.alignment = Alignment(horizontal="center")

            if is_active:
                active_rows += 1
                for j in range(1, len(HEADERS) + 1):
                    ws.cell(row=ri, column=j).fill = ACTIVE_FILL

    for col, w in {1: 52, 2: 14, 3: 26, 4: 12, 5: 12, 6: 22, 7: 15,
                   8: 15, 9: 15, 10: 14, 11: 14, 12: 15, 13: 14, 14: 38}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(row=4, column=1)
    ws.auto_filter.ref = f"A3:{get_column_letter(len(HEADERS))}{ri}"

    # ------------------------------------------------------------ methodology
    ms = wb.create_sheet("Methodology")
    ms.column_dimensions["A"].width = 34
    ms.column_dimensions["B"].width = 96
    ms.cell(row=1, column=1, value="Methodology & Data Coverage").font = \
        Font(bold=True, size=14, color=NAVY)

    def put(a, b, bold=False):
        row = ms.max_row + 1
        ms.cell(row=row, column=1, value=a).font = Font(bold=True, size=10)
        c = ms.cell(row=row, column=2, value=b)
        c.font = Font(bold=bold, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ms.append([])
    put("Source file", SRC)
    put("As-of date", f"{asof:%d %b %Y} (latest 'To Date' in the source file)")
    put("Stint rows listed", f"{len(kept):,}, across {len(order):,} schemes")
    put("Excluded", f"{len(passive):,} index funds / ETFs / fund-of-funds "
                    f"({dropped_rows:,} stint rows). They track an index instead of being "
                    f"actively managed, so a quartile rank against active equity peers says "
                    f"nothing about the manager. This removes rows only: peer groups and "
                    f"quartile cut-offs are unchanged, so figures match index.html.")
    put("Active manager rows", f"{active_rows:,} - highlighted green and bold; a stint is "
                               f"Active when its To Date equals the as-of date")
    ms.append([])

    put("Tenure", "To Date minus From Date, as 'X years and Y months'. Each row is one stint, "
                  "so a manager who left and returned has one row per spell. "
                  "'Tenure (Months)' is the same figure as an integer for sorting.")
    ms.append([])

    put("Rolling windows", f"{len(windows)} half-yearly windows, each exactly 3 years, running "
                           f"from '{windows[0][2]}' through '{windows[-1][2]}'. They overlap, so "
                           f"one long unbroken stint is counted in many of them.")
    put("Manager - Windows Managed",
        "Counts a window only when this manager held the fund for the WHOLE window (From Date "
        "on or before the window start AND To Date on or after the window end) and the fund had "
        "a ranked 3-year return for it. Partial windows are never counted.")
    put("Manager - Times in Top 25%",
        "Of the windows counted above, how many the fund finished in the top quartile of its "
        "peer group. Peer group = same sub-category, Regular plan, Growth option, funds with a "
        "return for that window; ranked by 3Y CAGR descending; the top max(1, round(n x 0.25)) "
        "are the top quartile. This is the rule index.html uses.")
    ms.append([])

    put("Fund - Windows Ranked",
        "Every window the FUND itself has a 3-year return for - in effect from its own start "
        "date through to the present - regardless of who was managing it. A fund younger than "
        "the full span is ranked in fewer than "
        f"{len(windows)} windows, so read the next two columns against this one.")
    put("Fund - Times in Top 25%",
        "Of those windows, how many the fund finished in the top quartile of its peer group. "
        "This is the fund's whole track record and is identical on every row of a given scheme; "
        "it is deliberately manager-independent, so compare it with the Manager columns to see "
        "how much of the record belongs to the current manager.")
    put("Fund - Top 25% Hit Rate", "Fund - Times in Top 25% divided by Fund - Windows Ranked.")
    ms.append([])

    put("Ties at the cut-off",
        "When two funds post an identical CAGR on the top-quartile boundary, peer-group order "
        "decides which takes the last slot, exactly as index.html does. One real case exists "
        "here: in Jan 2020 to Jan 2023 both Invesco India Mid Cap and Mirae Asset Midcap "
        "returned 24.98% for the 5th and last slot of a 20-fund Mid Cap group.")
    put("Attribution caveat",
        "A 3-year return is earned across the whole window, so crediting it to whoever held the "
        "fund throughout is an approximation: co-managers are not separated, and the source "
        "file lists each named manager independently. Where two managers ran the same fund over "
        "the same window, both rows show that window.", bold=True)
    ms.append([])

    put("Coverage - matched", f"{stats['matched rows']:,} listed stint rows matched a fund in the "
                              f"engine's Equity / Regular / Growth universe ({len(funds)} funds).")
    put("Coverage - no figures", f"{stats['unmatched rows']:,} listed stint rows have no window "
                                 f"figures; the Coverage column gives the reason per row. These "
                                 f"are long-short / SIF schemes with no 3-year history, and "
                                 f"sectoral funds missing from the engine's precomputed data.")
    ms.append([])
    put("Window list", "")
    for i, (s, e, lbl) in enumerate(windows, 1):
        ms.append([f"  Window {i}", f"{lbl}   ({s:%Y-%m-%d} to {e:%Y-%m-%d})"])

    target = OUT
    try:
        wb.save(target)
    except PermissionError:
        # the workbook is open in Excel, which holds an exclusive lock on it
        stem, ext = OUT.rsplit(".", 1)
        target = f"{stem}_new.{ext}"
        wb.save(target)
        print(f"\n!! {OUT} is open in Excel and could not be replaced.")
        print(f"!! Wrote {target} instead - close Excel and rerun, or keep this file.")

    print(f"\nwrote {target}  ({ri:,} sheet rows)")
    print(f"  matched rows           : {stats['matched rows']:,}")
    print(f"  unmatched rows         : {stats['unmatched rows']:,}")
    print(f"  rows with >=1 window   : {stats['rows with >=1 full window']:,}")
    print(f"  active rows highlighted: {active_rows:,}")


if __name__ == "__main__":
    main()
